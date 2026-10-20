from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import core


def frame() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "序号": [1, 2, 3, 4, 5, 6, 7, 8],
            "药品名称": [
                "阿司匹林",
                "阿司匹林",
                "阿司匹林",
                "维生素C",
                "维生素C",
                "布洛芬",
                " 布洛芬 ",
                "",
            ],
            "药品规格": ["100mg", "100mg", "50mg", "0.1g", "0.2g", "0.3g", "0.3g", "500ml"],
        }
    )


class CoreTests(unittest.TestCase):
    def test_simplify(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "清单.xlsx"
            frame().to_excel(path, index=False)
            sheet = core.load_workbook(path)
            self.assertTrue(sheet.usable)

            result = core.analyze(sheet, core.MODE_SIMPLIFY)
            self.assertEqual(result.total_rows, 8)
            self.assertEqual(result.valid_rows, 7)
            self.assertEqual(result.skipped_rows, 1)
            self.assertEqual(result.output_rows, 3)
            self.assertEqual(result.duplicate_groups, 3)
            self.assertEqual(result.removed_rows, 4)

            rows = result.df.to_dict("records")
            aspirin = next(row for row in rows if row["药品名称"] == "阿司匹林")
            self.assertEqual(aspirin["药品规格"], "100mg, 50mg")
            self.assertEqual(aspirin["重复次数"], 3)
            self.assertEqual(aspirin["原始序号"], "1, 2, 3")

            ibuprofen = next(row for row in rows if row["药品名称"] == "布洛芬")
            self.assertEqual(ibuprofen["重复次数"], 2)
            self.assertEqual(ibuprofen["原始序号"], "6, 7")

    def test_count(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "清单.xlsx"
            frame().to_excel(path, index=False)
            sheet = core.load_workbook(path)
            result = core.analyze(sheet, core.MODE_COUNT)
            self.assertEqual(result.output_rows, 5)
            self.assertEqual(result.duplicate_groups, 2)
            self.assertEqual(result.removed_rows, 2)

            rows = result.df.to_dict("records")
            self.assertEqual(rows[0]["序号"], 1)
            self.assertEqual(rows[0]["药品名称"], "阿司匹林")
            self.assertEqual(rows[0]["药品规格"], "100mg")
            self.assertEqual(rows[0]["重复次数"], 2)
            self.assertEqual(rows[0]["原来的序号位置"], "1, 2")
            self.assertEqual(rows[-1]["药品名称"], "布洛芬")
            self.assertEqual(rows[-1]["序号"], 6)

    def test_missing_columns(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "错误清单.xlsx"
            pd.DataFrame({"编号": [1], "名称": ["A"]}).to_excel(path, index=False)
            sheet = core.load_workbook(path)
            self.assertFalse(sheet.usable)
            self.assertIn("缺少必要列", sheet.problem or "")
            with self.assertRaises(core.UserError):
                core.analyze(sheet, core.MODE_SIMPLIFY)

    def test_empty_names(self) -> None:
        df = pd.DataFrame(
            {"序号": [1, 2], "药品名称": ["", "  "], "药品规格": ["A", "B"]}
        )
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "空名称.xlsx"
            df.to_excel(path, index=False)
            sheet = core.load_workbook(path)
            self.assertTrue(sheet.usable)
            with self.assertRaises(core.UserError):
                core.analyze(sheet, core.MODE_SIMPLIFY)

    def test_export_and_unique_names(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "清单.xlsx"
            frame().to_excel(path, index=False)
            sheet = core.load_workbook(path)
            result = core.analyze(sheet, core.MODE_SIMPLIFY)

            first = core.export_result(result, path)
            second = core.export_result(result, path)
            self.assertTrue(first.exists())
            self.assertTrue(second.exists())
            self.assertNotEqual(first, second)
            self.assertEqual(first.suffix, ".xlsx")

            wb = load_workbook(first)
            ws = wb.active
            self.assertEqual(ws.title, "精简结果")
            self.assertEqual(ws.freeze_panes, "A2")
            self.assertIsNotNone(ws.auto_filter.ref)
            self.assertEqual(ws["A1"].value, "药品名称")
            self.assertTrue(ws["A1"].font.bold)
            self.assertEqual(ws["A1"].fill.fgColor.rgb, "000F6CBD")
            self.assertFalse(ws.sheet_view.showGridLines)

    def test_xls_output_extension(self) -> None:
        # 仅验证命名函数会在 .xls 输入时输出 .xlsx，不依赖 xlwt。
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "旧清单.xls"
            source.write_bytes(b"")
            target = core._unique_output_path(source, core.MODE_COUNT, Path(tmp))
            self.assertEqual(target.suffix, ".xlsx")

    def test_header_aliases_and_details(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "别名表头.xlsx"
            data = [
                [None, None, None, None],
                ["编号", "药品名", "包装规格", "生产厂家"],
                [1, "阿莫西林胶囊", "0.25g", "甲厂"],
                [2, "阿莫西林胶囊", "0.25g", "甲厂"],
                [3, "阿莫西林颗粒", "0.25g", "乙厂"],
            ]
            pd.DataFrame(data).to_excel(path, index=False, header=False)
            sheet = core.load_workbook(path)
            self.assertTrue(sheet.usable)
            self.assertEqual(sheet.header_row, 1)
            self.assertEqual(sheet.column_map["药品名称"], 1)
            self.assertEqual(sheet.row_numbers, [3, 4, 5])
            self.assertEqual(sheet.df.columns[3], "生产厂家")

            result = core.analyze(sheet, core.MODE_COUNT)
            self.assertEqual(result.output_rows, 2)
            self.assertEqual(len(result.details), 2)
            self.assertEqual(result.details.iloc[0]["原始行号"], 3)
            self.assertEqual(result.details.iloc[0]["生产厂家"], "甲厂")
            self.assertIn("源文件", result.summary)

    def test_manual_column_mapping(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "手动映射.xlsx"
            pd.DataFrame(
                {"编号": [1, 2], "药名": ["药A", "药A"], "包装": ["1盒", "1盒"]}
            ).to_excel(path, index=False)
            auto = core.load_workbook(path)
            self.assertFalse(auto.usable)
            self.assertIn("药品规格", auto.missing_columns)
            mapped = core.load_workbook(
                path,
                header_row=auto.header_row,
                column_map={"序号": 0, "药品名称": 1, "药品规格": 2},
            )
            self.assertTrue(mapped.usable)
            self.assertEqual(mapped.total_rows, 2)
            result = core.analyze(mapped, core.MODE_SIMPLIFY)
            self.assertEqual(result.output_rows, 1)
            self.assertEqual(result.df.iloc[0]["重复次数"], 2)

    def test_normalisation_and_name_aliases(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "规范药名.xlsx"
            pd.DataFrame(
                {
                    "序号": [1, 2, 3],
                    "药品名称": [
                        "阿莫西林胶囊（0.25g）",
                        "阿莫西林胶囊(0.25g)",
                        "阿莫西林胶囊",
                    ],
                    "药品规格": ["A", "A", "B"],
                }
            ).to_excel(path, index=False)
            sheet = core.load_workbook(path)
            result = core.analyze(sheet, core.MODE_SIMPLIFY)
            # 全角括号和半角括号属于安全的格式差异，自动归为一组。
            self.assertEqual(result.output_rows, 2)
            self.assertEqual(result.normalised_name_groups, 1)
            details = result.details
            self.assertEqual(len(details), 2)
            self.assertEqual(details.iloc[0]["分组药名"], "阿莫西林胶囊（0.25g）")

    def test_name_candidates_and_alias_merge(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "疑似药名.xlsx"
            pd.DataFrame(
                {
                    "序号": [1, 2, 3],
                    "药品名称": ["阿莫西林胶囊", "阿莫西林颗粒", "维生素C片"],
                    "药品规格": ["A", "B", "C"],
                }
            ).to_excel(path, index=False)
            sheet = core.load_workbook(path)
            candidates = core.find_name_candidates(sheet)
            self.assertTrue(candidates)
            names = {variant.display for variant in candidates[0].variants}
            self.assertIn("阿莫西林胶囊", names)
            self.assertIn("阿莫西林颗粒", names)
            self.assertTrue(any(variant.specs for variant in candidates[0].variants))

            canonical = max(candidates[0].variants, key=lambda v: (v.count, len(v.key)))
            aliases = {
                variant.key: canonical.key
                for variant in candidates[0].variants
                if variant.key != canonical.key
            }
            merged = core.analyze(sheet, core.MODE_SIMPLIFY, name_aliases=aliases)
            self.assertEqual(merged.output_rows, 2)
            self.assertEqual(merged.alias_groups, 1)

    def test_sort_orders(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "排序.xlsx"
            frame().to_excel(path, index=False)
            sheet = core.load_workbook(path)
            by_first = core.analyze(sheet, core.MODE_SIMPLIFY, sort_order="first")
            by_name = core.analyze(sheet, core.MODE_SIMPLIFY, sort_order="name")
            by_count = core.analyze(sheet, core.MODE_SIMPLIFY, sort_order="count_desc")
            self.assertEqual(by_first.df.iloc[0]["药品名称"], "阿司匹林")
            self.assertEqual(
                by_name.df["药品名称"].tolist(),
                sorted(by_name.df["药品名称"].tolist()),
            )
            self.assertEqual(
                by_count.df["重复次数"].tolist(),
                sorted(by_count.df["重复次数"].tolist(), reverse=True),
            )

    def test_export_contains_detail_and_summary(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "带明细.xlsx"
            frame().to_excel(path, index=False)
            sheet = core.load_workbook(path)
            result = core.analyze(sheet, core.MODE_SIMPLIFY)
            output = core.export_result(
                result,
                path,
                app_version="1.1.0",
                output_label="Sheet 1",
            )
            self.assertIn("_Sheet 1_", output.name)
            wb = load_workbook(output)
            self.assertEqual(
                wb.sheetnames,
                ["精简结果", "重复明细", "处理汇总"],
            )
            self.assertEqual(wb["处理汇总"]["A1"].value, "项目")
            summary = {
                row[0]: row[1]
                for row in wb["处理汇总"].iter_rows(min_row=2, values_only=True)
            }
            self.assertEqual(summary["工具版本"], "1.1.0")
            self.assertEqual(summary["结果条数"], 3)


if __name__ == "__main__":
    unittest.main(verbosity=2)
