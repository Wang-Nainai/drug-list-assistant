"""核心数据处理：读取 Excel、药名核对、两种去重统计方式、导出带格式的 Excel。"""

from __future__ import annotations

import difflib
import math
import numbers
import os
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Iterable, Literal

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

Mode = Literal["simplify", "count"]

MODE_SIMPLIFY: Mode = "simplify"
MODE_COUNT: Mode = "count"

MODE_LABELS = {
    MODE_SIMPLIFY: "精简合并",
    MODE_COUNT: "精确统计",
}
OUTPUT_SUFFIX = {
    MODE_SIMPLIFY: "精简统计",
    MODE_COUNT: "重复统计",
}
RESULT_SHEET = {
    MODE_SIMPLIFY: "精简结果",
    MODE_COUNT: "统计结果",
}
DETAIL_SHEET = "重复明细"
SUMMARY_SHEET = "处理汇总"

SORT_ORDER_LABELS = {
    "first": "首次出现",
    "name": "药名排序",
    "count_desc": "重复次数降序",
}

REQUIRED_COLUMNS: tuple[str, ...] = ("序号", "药品名称", "药品规格")
SUPPORTED_EXTENSIONS = {".xlsx", ".xls", ".xlsm"}

SIMPLIFY_COLUMNS = ["药品名称", "药品规格", "重复次数", "原始序号"]
COUNT_COLUMNS = ["序号", "药品名称", "药品规格", "重复次数", "原来的序号位置"]

DETAIL_BASE_COLUMNS = [
    "分组药名",
    "分组规格",
    "原始行号",
    "序号",
    "药品名称",
    "药品规格",
]
DETAIL_OPTIONAL_COLUMNS = (
    "生产厂家",
    "交易时间",
    "交易类型名称",
    "交易单位",
    "交易数量",
    "交易金额",
    "批号",
    "有效期",
)

COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "序号": ("序号", "编号", "序", "no", "no.", "id"),
    "药品名称": ("药品名称", "药品名", "药名", "名称", "品名"),
    "药品规格": ("药品规格", "规格", "包装规格", "规格型号"),
}

HEADER_SCAN_ROWS = 30
MAX_CANDIDATE_NAMES = 2500
DEFAULT_CANDIDATE_THRESHOLD = 0.74
DOSAGE_SUFFIXES = (
    "干混悬剂",
    "吸入气雾剂",
    "缓释胶囊",
    "缓释片",
    "控释片",
    "分散片",
    "肠溶片",
    "咀嚼片",
    "泡腾片",
    "口服溶液",
    "口服液",
    "注射液",
    "滴眼液",
    "滴耳液",
    "喷雾剂",
    "气雾剂",
    "混悬液",
    "乳膏",
    "软膏",
    "凝胶",
    "糖浆",
    "胶囊",
    "颗粒",
    "滴丸",
    "片",
    "丸",
    "散",
    "贴",
    "栓",
    "滴剂",
)


class UserError(Exception):
    """可以直接展示给用户的错误。"""


@dataclass
class HeaderLayout:
    header_row: int
    columns: list[str]
    mapping: dict[str, int]
    raw_rows: list[list[str]] = field(default_factory=list)

    @property
    def missing(self) -> tuple[str, ...]:
        return tuple(column for column in REQUIRED_COLUMNS if column not in self.mapping)


@dataclass
class SheetData:
    path: Path
    sheet_names: list[str]
    sheet_name: str
    df: pd.DataFrame
    missing_columns: tuple[str, ...] = ()
    duplicated_required_columns: tuple[str, ...] = ()
    problem: str | None = None
    header_row: int = 0
    raw_columns: list[str] = field(default_factory=list)
    column_map: dict[str, int] = field(default_factory=dict)
    row_numbers: list[int] = field(default_factory=list)

    @property
    def usable(self) -> bool:
        return self.problem is None

    @property
    def total_rows(self) -> int:
        return int(len(self.df.index))

    @property
    def total_columns(self) -> int:
        return int(len(self.df.columns))


@dataclass
class NameVariant:
    key: str
    display: str
    count: int
    specs: list[str] = field(default_factory=list)


@dataclass
class NameCandidateGroup:
    variants: list[NameVariant]
    score: float = 0.0

    @property
    def total_count(self) -> int:
        return sum(variant.count for variant in self.variants)


@dataclass
class AnalysisResult:
    mode: Mode
    df: pd.DataFrame
    total_rows: int
    valid_rows: int
    skipped_rows: int
    output_rows: int
    duplicate_groups: int
    removed_rows: int
    columns: list[str] = field(default_factory=list)
    details: pd.DataFrame = field(default_factory=pd.DataFrame)
    summary: dict[str, object] = field(default_factory=dict)
    source_file: str = ""
    source_sheet: str = ""
    alias_groups: int = 0
    normalised_name_groups: int = 0

    def __post_init__(self) -> None:
        if not self.columns:
            self.columns = [str(c) for c in self.df.columns]


def validate_path(file_path: str | os.PathLike[str]) -> Path:
    path = Path(file_path).expanduser()
    if not path.exists():
        raise UserError(f"找不到文件：{path.name or path}")
    if not path.is_file():
        raise UserError("请选择一个 Excel 文件，而不是文件夹。")
    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise UserError("仅支持 .xlsx、.xls 和 .xlsm 格式的 Excel 文件。")
    return path


def _header_key(value: object) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKC", text).replace("\ufeff", "")
    text = re.sub(r"\s+", "", text).strip("*:：")
    return text.casefold()


def _match_required_column(value: object) -> str | None:
    key = _header_key(value)
    if not key:
        return None
    for canonical, aliases in COLUMN_ALIASES.items():
        for alias in aliases:
            alias_key = _header_key(alias)
            if key == alias_key or key.startswith(alias_key):
                return canonical
    return None


def _normalise_column_name(value: object, index: int) -> str:
    if value is None:
        return f"列{index + 1}"
    text = unicodedata.normalize("NFKC", str(value)).replace("\ufeff", "").strip()
    return text or f"列{index + 1}"


def _normalise_frame(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy(deep=False)
    df.columns = [_normalise_column_name(c, i) for i, c in enumerate(df.columns)]
    return df


def _map_header_row(values: Iterable[object]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for index, value in enumerate(values):
        canonical = _match_required_column(value)
        if canonical and canonical not in mapping:
            mapping[canonical] = index
    return mapping


def _preview_layout(excel_file: pd.ExcelFile, sheet_name: str) -> HeaderLayout:
    try:
        preview = excel_file.parse(
            sheet_name=sheet_name,
            header=None,
            nrows=HEADER_SCAN_ROWS,
            dtype=object,
        )
    except Exception:
        return HeaderLayout(0, [], {}, [])

    best_row = 0
    best_mapping: dict[str, int] = {}
    best_score = -1
    raw_rows: list[list[str]] = []
    for row_index, row in enumerate(preview.itertuples(index=False, name=None)):
        values = ["" if value is None else str(value) for value in row]
        raw_rows.append(values)
        mapping = _map_header_row(row)
        score = len(mapping)
        if score > best_score:
            best_row = row_index
            best_mapping = mapping
            best_score = score

    if best_score < 2:
        best_row = 0
        if raw_rows:
            best_mapping = _map_header_row(raw_rows[0])
        else:
            best_mapping = {}

    columns = (
        [_normalise_column_name(value, index) for index, value in enumerate(raw_rows[best_row])]
        if raw_rows and best_row < len(raw_rows)
        else []
    )
    return HeaderLayout(best_row, columns, best_mapping, raw_rows)


def _preview_raw_rows(raw: pd.DataFrame) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in raw.itertuples(index=False, name=None):
        rows.append(["" if value is None else str(value) for value in row])
    return rows


def inspect_columns(file_path: str | os.PathLike[str], sheet_name: str | None = None) -> HeaderLayout:
    """只读取表头附近的内容，用于自动识别和手动列映射。"""
    path = validate_path(file_path)
    try:
        excel_file = pd.ExcelFile(path)
    except Exception as exc:  # noqa: BLE001
        raise UserError("无法读取这个 Excel 文件。") from exc
    try:
        names = [str(name) for name in excel_file.sheet_names]
        selected = sheet_name if sheet_name is not None else names[0]
        if selected not in names:
            raise UserError(f"找不到工作表「{selected}」。")
        return _preview_layout(excel_file, selected)
    finally:
        try:
            excel_file.close()
        except Exception:  # noqa: BLE001
            pass


def load_workbook(
    file_path: str | os.PathLike[str],
    sheet_name: str | None = None,
    *,
    header_row: int | None = None,
    column_map: dict[str, int] | None = None,
) -> SheetData:
    """读取指定工作表，自动识别表头，也支持手动指定列位置。"""
    path = validate_path(file_path)

    try:
        excel_file = pd.ExcelFile(path)
    except ImportError as exc:
        raise UserError("当前环境缺少读取该 Excel 格式所需的组件。") from exc
    except Exception as exc:  # noqa: BLE001 - 统一转换为友好提示
        raise UserError(
            "无法读取这个 Excel 文件。文件可能已损坏、被加密，或格式不受支持。"
        ) from exc

    try:
        sheet_names = [str(name) for name in excel_file.sheet_names]
        if not sheet_names:
            raise UserError("这个 Excel 文件中没有工作表。")

        selected = sheet_name if sheet_name is not None else sheet_names[0]
        if selected not in sheet_names:
            raise UserError(f"找不到工作表「{selected}」。")

        layout = _preview_layout(excel_file, selected)
        selected_header = layout.header_row if header_row is None else header_row
        if selected_header < 0:
            selected_header = 0

        try:
            df = excel_file.parse(
                sheet_name=selected,
                header=selected_header,
                dtype=object,
            )
        except Exception as exc:  # noqa: BLE001
            raise UserError(f"无法读取工作表「{selected}」，请检查该表内容。") from exc
    finally:
        try:
            excel_file.close()
        except Exception:  # noqa: BLE001
            pass

    raw_columns = [_normalise_column_name(c, i) for i, c in enumerate(df.columns)]
    if column_map is None:
        if header_row is None:
            mapping = dict(layout.mapping)
        else:
            mapping = _map_header_row(df.columns)
    else:
        mapping = {
            canonical: int(index)
            for canonical, index in column_map.items()
            if canonical in REQUIRED_COLUMNS and int(index) >= 0
        }

    df = df.copy(deep=False)
    if len(df.columns):
        new_columns = list(df.columns)
        for canonical, index in mapping.items():
            if index < len(new_columns):
                new_columns[index] = canonical
        df.columns = new_columns
    df = _normalise_frame(df)

    row_numbers = [int(index) + selected_header + 2 for index in df.index]
    keep = ~df.isna().all(axis=1)
    df = df.loc[keep]
    row_numbers = [row for row, keep_row in zip(row_numbers, keep.tolist()) if keep_row]
    df = df.reset_index(drop=True)

    missing = tuple(col for col in REQUIRED_COLUMNS if col not in df.columns)
    duplicated = tuple(
        col for col in REQUIRED_COLUMNS if list(df.columns).count(col) > 1
    )

    problem: str | None = None
    if df.empty or len(df.columns) == 0:
        problem = "当前工作表没有数据。"
    elif missing:
        problem = "缺少必要列：" + "、".join(missing) + "。"
    elif duplicated:
        problem = "表头中有重复的必要列：" + "、".join(duplicated) + "。"

    return SheetData(
        path=path,
        sheet_names=sheet_names,
        sheet_name=selected,
        df=df,
        missing_columns=missing,
        duplicated_required_columns=duplicated,
        problem=problem,
        header_row=selected_header,
        raw_columns=raw_columns,
        column_map=mapping,
        row_numbers=row_numbers,
    )


def load_sheet(
    file_path: str | os.PathLike[str],
    sheet_name: str,
    *,
    header_row: int | None = None,
    column_map: dict[str, int] | None = None,
) -> SheetData:
    return load_workbook(
        file_path,
        sheet_name,
        header_row=header_row,
        column_map=column_map,
    )


def normalise_match_text(value: object) -> str:
    """把名称/规格转换成用于比较的文本，不改动最终展示的原值。"""
    text = _clean_text(value)
    if not text:
        return ""
    text = unicodedata.normalize("NFKC", text)
    text = (
        text.replace("\u200b", "")
        .replace("\u200c", "")
        .replace("\u200d", "")
        .replace("\ufeff", "")
    )
    text = re.sub(r"\s+", " ", text).strip()
    text = re.sub(r"\s*([()\[\]{}【】,，、:：;；])\s*", r"\1", text)
    return text


def _resolve_alias(key: str, aliases: dict[str, str]) -> str:
    seen: set[str] = set()
    current = key
    while current in aliases and aliases[current] != current and current not in seen:
        seen.add(current)
        current = aliases[current]
    return current


def _seq_text(value: object) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass

    if isinstance(value, pd.Timestamp):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S") if value.time() else value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, numbers.Integral):
        return str(int(value))
    if isinstance(value, numbers.Real):
        number = float(value)
        if not math.isfinite(number):
            return ""
        if number.is_integer():
            return str(int(number))
        return f"{number:g}"

    text = str(value).strip()
    if text.lower() in {"nan", "nat", "none"}:
        return ""
    return text


def _seq_value(value: object) -> object:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass

    if isinstance(value, numbers.Integral):
        return int(value)
    if isinstance(value, numbers.Real):
        number = float(value)
        if math.isfinite(number) and number.is_integer():
            return int(number)
    return _seq_text(value)


def _clean_text(value: object) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    return str(value).strip()


def _detail_value(value: object) -> object:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(value, pd.Timestamp):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S") if value.time() else value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, numbers.Integral):
        return int(value)
    if isinstance(value, numbers.Real):
        number = float(value)
        return int(number) if math.isfinite(number) and number.is_integer() else number
    return str(value).strip()


def _join_specs(specs: list[str]) -> str:
    if any("," in spec or "，" in spec for spec in specs):
        return "；".join(specs)
    return ", ".join(specs)


def _name_base(key: str) -> str:
    """仅供候选提示使用：去掉常见剂型后缀和括号内容。"""
    text = re.sub(r"\([^)]*\)", "", key).strip()
    changed = True
    while changed and text:
        changed = False
        for suffix in DOSAGE_SUFFIXES:
            if text.endswith(suffix) and len(text) > len(suffix) + 1:
                text = text[: -len(suffix)].strip()
                changed = True
                break
    return text


def find_name_candidates(
    sheet: SheetData,
    name_aliases: dict[str, str] | None = None,
    *,
    threshold: float = DEFAULT_CANDIDATE_THRESHOLD,
    max_groups: int = 80,
) -> list[NameCandidateGroup]:
    """找出写法相近、但仍未合并的药名，交给用户人工确认。"""
    if not sheet.usable:
        return []

    aliases = {
        normalise_match_text(key): normalise_match_text(value)
        for key, value in (name_aliases or {}).items()
        if normalise_match_text(key) and normalise_match_text(value)
    }
    grouped: dict[str, dict[str, object]] = {}
    for value, spec_value in zip(
        sheet.df["药品名称"].tolist(), sheet.df["药品规格"].tolist()
    ):
        display = _clean_text(value)
        key = normalise_match_text(display)
        if not key:
            continue
        canonical = _resolve_alias(key, aliases)
        entry = grouped.get(canonical)
        if entry is None:
            entry = {
                "key": canonical,
                "display": display,
                "count": 0,
                "specs": [],
            }
            grouped[canonical] = entry
        entry["count"] = int(entry["count"]) + 1
        spec = _clean_text(spec_value)
        specs = entry["specs"]
        assert isinstance(specs, list)
        if spec and spec not in specs and len(specs) < 6:
            specs.append(spec)

    variants = [
        NameVariant(
            str(item["key"]),
            str(item["display"]),
            int(item["count"]),
            list(item["specs"]),
        )
        for item in grouped.values()
    ]
    if len(variants) > MAX_CANDIDATE_NAMES:
        variants = sorted(variants, key=lambda v: v.count, reverse=True)[:MAX_CANDIDATE_NAMES]
    variants.sort(key=lambda v: v.count, reverse=True)

    parent = list(range(len(variants)))

    def find(index: int) -> int:
        while parent[index] != index:
            parent[index] = parent[parent[index]]
            index = parent[index]
        return index

    def union(left: int, right: int) -> None:
        left_root, right_root = find(left), find(right)
        if left_root != right_root:
            parent[right_root] = left_root

    scores: dict[tuple[int, int], float] = {}
    bases = [_name_base(variant.key) for variant in variants]
    for left in range(len(variants)):
        a = variants[left].key
        for right in range(left + 1, len(variants)):
            b = variants[right].key
            base_a, base_b = bases[left], bases[right]
            same_base = (
                bool(base_a)
                and base_a == base_b
                and len(base_a) >= 2
                and a != b
            )
            length_gap = abs(len(a) - len(b))
            if not same_base and length_gap > max(3, int(max(len(a), len(b)) * 0.45)):
                continue
            ratio = difflib.SequenceMatcher(None, a, b, autojunk=False).ratio()
            if ratio >= threshold or same_base:
                union(left, right)
                scores[(left, right)] = max(ratio, 0.75 if same_base else 0.0)

    components: dict[int, list[int]] = {}
    for index in range(len(variants)):
        components.setdefault(find(index), []).append(index)

    groups: list[NameCandidateGroup] = []
    for indexes in components.values():
        if len(indexes) < 2:
            continue
        group_score = 0.0
        for i, left in enumerate(indexes):
            for right in indexes[i + 1 :]:
                group_score = max(
                    group_score,
                    scores.get((min(left, right), max(left, right)), 0.0),
                )
        groups.append(
            NameCandidateGroup(
                variants=[variants[index] for index in indexes],
                score=group_score,
            )
        )

    groups.sort(key=lambda group: (group.total_count, group.score), reverse=True)
    return groups[:max_groups]


def analyze(
    sheet: SheetData,
    mode: Mode,
    *,
    name_aliases: dict[str, str] | None = None,
    sort_order: str = "first",
) -> AnalysisResult:
    """按指定方式生成结果，保持药品首次出现的顺序或按指定方式排序。"""
    if mode not in (MODE_SIMPLIFY, MODE_COUNT):
        raise ValueError(f"未知的处理方式：{mode}")
    if not sheet.usable:
        raise UserError(sheet.problem or "当前工作表不可用。")
    if sort_order not in SORT_ORDER_LABELS:
        sort_order = "first"

    aliases = {
        normalise_match_text(key): normalise_match_text(value)
        for key, value in (name_aliases or {}).items()
        if normalise_match_text(key) and normalise_match_text(value)
    }
    alias_groups = len({value for key, value in aliases.items() if key != value})

    df = sheet.df
    original_names = [_clean_text(v) for v in df["药品名称"].tolist()]
    original_specs = [_clean_text(v) for v in df["药品规格"].tolist()]
    names = [normalise_match_text(value) for value in original_names]
    specs = [normalise_match_text(value) for value in original_specs]
    seqs = df["序号"].tolist()

    total_rows = int(len(df.index))
    valid_rows = sum(1 for name in names if name)
    skipped_rows = total_rows - valid_rows
    if valid_rows == 0:
        raise UserError("没有找到有效的药品数据：药品名称列全部为空。")

    display_by_key: dict[str, str] = {}
    for original, key in zip(original_names, names):
        if key and key not in display_by_key:
            display_by_key[key] = original
    normalised_forms: dict[str, set[str]] = {}
    for original, key in zip(original_names, names):
        if key:
            normalised_forms.setdefault(key, set()).add(original)
    normalised_name_groups = sum(
        1 for forms in normalised_forms.values() if len(forms) > 1
    )

    optional_columns = [
        column for column in DETAIL_OPTIONAL_COLUMNS if column in df.columns
    ]
    optional_values = {
        column: [_detail_value(value) for value in df[column].tolist()]
        for column in optional_columns
    }
    row_numbers = (
        list(sheet.row_numbers)
        if len(sheet.row_numbers) == total_rows
        else [index + sheet.header_row + 2 for index in range(total_rows)]
    )

    detail_rows: list[dict[str, object]] = []

    if mode == MODE_SIMPLIFY:
        groups: dict[str, dict[str, object]] = {}
        for index, (original_name, name, original_spec, spec, seq) in enumerate(
            zip(original_names, names, original_specs, specs, seqs)
        ):
            if not name:
                continue
            canonical = _resolve_alias(name, aliases)
            group = groups.get(canonical)
            if group is None:
                display = (
                    display_by_key.get(canonical)
                    or display_by_key.get(name)
                    or original_name
                )
                group = {
                    "display": display,
                    "specs": [],
                    "seen": set(),
                    "positions": [],
                    "count": 0,
                    "rows": [],
                }
                groups[canonical] = group

            group["count"] = int(group["count"]) + 1
            if spec:
                seen = group["seen"]
                assert isinstance(seen, set)
                if spec not in seen:
                    seen.add(spec)
                    specs_list = group["specs"]
                    assert isinstance(specs_list, list)
                    specs_list.append(original_spec or spec)

            position = _seq_text(seq)
            if position:
                positions = group["positions"]
                assert isinstance(positions, list)
                positions.append(position)

            detail = {
                "分组药名": group["display"],
                "原始行号": row_numbers[index],
                "序号": _detail_value(seq),
                "药品名称": original_name,
                "药品规格": original_spec,
            }
            for column in optional_columns:
                detail[column] = optional_values[column][index]
            group["rows"].append(detail)  # type: ignore[union-attr]

        records: list[dict[str, object]] = []
        for group in groups.values():
            spec_text = _join_specs(list(group["specs"]))  # type: ignore[arg-type]
            records.append(
                {
                    "药品名称": group["display"],
                    "药品规格": spec_text,
                    "重复次数": group["count"],
                    "原始序号": ", ".join(group["positions"]),  # type: ignore[arg-type]
                }
            )
            if int(group["count"]) > 1:
                for detail in group["rows"]:  # type: ignore[union-attr]
                    detail["分组规格"] = spec_text
                    detail_rows.append(detail)
        output = pd.DataFrame(records, columns=SIMPLIFY_COLUMNS)
    else:
        entries: dict[tuple[str, str], dict[str, object]] = {}
        for index, (original_name, name, original_spec, spec, seq) in enumerate(
            zip(original_names, names, original_specs, specs, seqs)
        ):
            if not name:
                continue
            canonical = _resolve_alias(name, aliases)
            key = (canonical, spec)
            entry = entries.get(key)
            if entry is None:
                display_name = (
                    display_by_key.get(canonical)
                    or display_by_key.get(name)
                    or original_name
                )
                entry = {
                    "序号": _seq_value(seq),
                    "药品名称": display_name,
                    "药品规格": original_spec or spec,
                    "重复次数": 0,
                    "positions": [],
                    "rows": [],
                }
                entries[key] = entry

            entry["重复次数"] = int(entry["重复次数"]) + 1
            position = _seq_text(seq)
            if position:
                positions = entry["positions"]
                assert isinstance(positions, list)
                positions.append(position)

            detail = {
                "分组药名": entry["药品名称"],
                "分组规格": entry["药品规格"],
                "原始行号": row_numbers[index],
                "序号": _detail_value(seq),
                "药品名称": original_name,
                "药品规格": original_spec,
            }
            for column in optional_columns:
                detail[column] = optional_values[column][index]
            entry["rows"].append(detail)  # type: ignore[union-attr]

        records = []
        for entry in entries.values():
            records.append(
                {
                    "序号": entry["序号"],
                    "药品名称": entry["药品名称"],
                    "药品规格": entry["药品规格"],
                    "重复次数": entry["重复次数"],
                    "原来的序号位置": ", ".join(entry["positions"]),  # type: ignore[arg-type]
                }
            )
            if int(entry["重复次数"]) > 1:
                detail_rows.extend(entry["rows"])  # type: ignore[arg-type]
        output = pd.DataFrame(records, columns=COUNT_COLUMNS)

    if sort_order == "name":
        output = output.sort_values(
            by="药品名称", key=lambda series: series.astype(str).str.casefold(), kind="stable"
        ).reset_index(drop=True)
    elif sort_order == "count_desc":
        output = output.sort_values(
            by="重复次数", ascending=False, kind="stable"
        ).reset_index(drop=True)

    columns = list(DETAIL_BASE_COLUMNS) + optional_columns
    details = pd.DataFrame(detail_rows, columns=columns)

    output_rows = int(len(output.index))
    duplicate_groups = int((output["重复次数"] > 1).sum()) if output_rows else 0

    summary = {
        "源文件": sheet.path.name,
        "工作表": sheet.sheet_name,
        "处理模式": MODE_LABELS[mode],
        "排序方式": SORT_ORDER_LABELS[sort_order],
        "原始行数": total_rows,
        "有效行数": valid_rows,
        "跳过空药名行": skipped_rows,
        "结果条数": output_rows,
        "重复组数": duplicate_groups,
        "可减少行": max(valid_rows - output_rows, 0),
        "重复明细行数": int(len(details)),
        "自动规范化药名组": normalised_name_groups,
        "手工合并药名组": alias_groups,
    }

    return AnalysisResult(
        mode=mode,
        df=output,
        total_rows=total_rows,
        valid_rows=valid_rows,
        skipped_rows=skipped_rows,
        output_rows=output_rows,
        duplicate_groups=duplicate_groups,
        removed_rows=max(valid_rows - output_rows, 0),
        columns=[str(c) for c in output.columns],
        details=details,
        summary=summary,
        source_file=sheet.path.name,
        source_sheet=sheet.sheet_name,
        alias_groups=alias_groups,
        normalised_name_groups=normalised_name_groups,
    )


def _safe_filename_part(value: str) -> str:
    text = re.sub(r"[\\/:*?\"<>|]+", "_", str(value)).strip(" .")
    return text[:40]


def _unique_output_path(
    source: Path,
    mode: Mode,
    output_dir: Path,
    label: str = "",
) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M")
    stem = source.stem[:80] or "药品清单"
    safe_label = _safe_filename_part(label)
    label_part = f"_{safe_label}" if safe_label else ""
    base = f"{stem}{label_part}_{OUTPUT_SUFFIX[mode]}_{timestamp}"
    candidate = output_dir / f"{base}.xlsx"
    counter = 2
    while candidate.exists():
        candidate = output_dir / f"{base}_{counter}.xlsx"
        counter += 1
    return candidate


def _display_width(value: object) -> int:
    text = "" if value is None else str(value)
    width = 0
    for char in text:
        width += 2 if ord(char) > 127 else 1
    return width


def _style_worksheet(ws, df: pd.DataFrame, *, auto_filter: bool = True) -> None:
    thin_side = Side(style="thin", color="E1E5EA")
    border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    header_fill = PatternFill("solid", fgColor="0F6CBD")
    band_fill = PatternFill("solid", fgColor="F7FAFD")
    header_font = Font(name="Microsoft YaHei", size=10.5, bold=True, color="FFFFFF")
    body_font = Font(name="Microsoft YaHei", size=10, color="1F1F1F")

    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "0F6CBD"
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    max_row = int(ws.max_row)
    max_col = int(ws.max_column)
    numeric_columns = {"序号", "重复次数", "原始行号"}
    headers = [str(c) for c in df.columns]

    style_body = max_row <= 20000
    if style_body:
        for row_idx in range(2, max_row + 1):
            ws.row_dimensions[row_idx].height = 22
            for col_idx in range(1, max_col + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                header = headers[col_idx - 1] if col_idx - 1 < len(headers) else ""
                cell.font = body_font
                cell.border = border
                cell.alignment = Alignment(
                    horizontal="center" if header in numeric_columns else "left",
                    vertical="center",
                    indent=0 if header in numeric_columns else 1,
                )
                if row_idx % 2 == 0:
                    cell.fill = band_fill

    for col_idx, header in enumerate(headers, start=1):
        sample = df.iloc[:1000, col_idx - 1] if col_idx - 1 < len(df.columns) else []
        content_width = max([_display_width(header), *(_display_width(v) for v in sample)])
        width = min(max(10, content_width + 4), 52)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
        if header in numeric_columns:
            for row_idx in range(2, max_row + 1):
                ws.cell(row=row_idx, column=col_idx).number_format = "0"

    if auto_filter and max_row >= 1 and max_col >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"
    ws.sheet_view.zoomScale = 100


def _write_frame(writer, sheet_name: str, df: pd.DataFrame, *, auto_filter: bool = True) -> None:
    df.to_excel(writer, sheet_name=sheet_name, index=False)
    _style_worksheet(writer.sheets[sheet_name], df, auto_filter=auto_filter)


def export_result(
    result: AnalysisResult,
    source_path: str | os.PathLike[str],
    output_dir: str | os.PathLike[str] | None = None,
    *,
    app_version: str = "",
    output_label: str = "",
) -> Path:
    source = Path(source_path)
    target_dir = Path(output_dir) if output_dir else source.parent
    try:
        target_dir.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise UserError(f"无法创建导出文件夹：{target_dir}") from exc

    output_path = _unique_output_path(
        source,
        result.mode,
        target_dir,
        label=output_label,
    )
    summary = dict(result.summary)
    summary["处理时间"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    if app_version:
        summary["工具版本"] = app_version
    summary_df = pd.DataFrame(list(summary.items()), columns=["项目", "内容"])

    try:
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            _write_frame(writer, RESULT_SHEET[result.mode], result.df)
            if not result.details.empty:
                _write_frame(writer, DETAIL_SHEET, result.details)
            _write_frame(writer, SUMMARY_SHEET, summary_df, auto_filter=False)
            workbook = writer.book
            workbook.properties.creator = "药单整理助手"
            workbook.properties.lastModifiedBy = "药单整理助手"
            workbook.properties.title = f"{source.stem} - {MODE_LABELS[result.mode]}"
            workbook.properties.subject = "药品清单去重与重复统计"
            workbook.properties.keywords = "药品清单, 去重, 重复统计, 明细, 汇总"
            for sheet in workbook.worksheets:
                if sheet.title == SUMMARY_SHEET:
                    sheet.sheet_properties.tabColor = "0E7A4E"
                elif sheet.title == DETAIL_SHEET:
                    sheet.sheet_properties.tabColor = "B25E00"
    except PermissionError as exc:
        _remove_partial(output_path)
        raise UserError(
            "无法写入结果文件，目标文件可能正在 Excel 中打开。请关闭后重试。"
        ) from exc
    except Exception as exc:  # noqa: BLE001
        _remove_partial(output_path)
        raise UserError(f"导出 Excel 时出错：{exc}") from exc

    return output_path


def _remove_partial(path: Path) -> None:
    try:
        if path.exists():
            path.unlink()
    except OSError:
        pass


def result_to_rows(result: AnalysisResult, limit: int | None = None) -> Iterable[list[str]]:
    df = result.df if limit is None else result.df.head(limit)
    for record in df.itertuples(index=False, name=None):
        yield [("" if v is None else str(v)) for v in record]
